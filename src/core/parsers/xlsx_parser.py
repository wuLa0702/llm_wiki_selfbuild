"""XLSX 解析器 — 提取文本为 Markdown 表格"""
from src.core.logging_config import get_logger

logger = get_logger("parsers.xlsx")


class XlsxParser:
    """XLSX 文件解析器

    每个 worksheet → `## Sheet 名` 标题。
    行提取为 Markdown 表格。
    """

    @staticmethod
    def parse(file_path: str) -> str:
        """解析 XLSX 文件为 Markdown 文本

        Args:
            file_path: .xlsx 文件路径

        Returns:
            Markdown 格式的结构化文本
        """
        from openpyxl import load_workbook

        wb = load_workbook(file_path, read_only=True, data_only=True)
        parts: list[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_data: list[list[str]] = []
            for row in ws.iter_rows():
                cells = [str(cell.value) if cell.value is not None else "" for cell in row]
                # 跳过全空行
                if any(c.strip() for c in cells):
                    rows_data.append(cells)

            if not rows_data:
                parts.append(f"## {sheet_name}\n\n（空表）")
                continue

            # 构建 Markdown 表格
            md_rows = []
            for i, row in enumerate(rows_data):
                md_rows.append("| " + " | ".join(row) + " |")
                if i == 0:
                    # 表头分隔行
                    md_rows.append("| " + " | ".join(["---"] * len(row)) + " |")

            parts.append(f"## {sheet_name}\n\n" + "\n".join(md_rows))

        wb.close()
        result = "\n\n".join(parts)
        logger.info("XLSX 解析完成 | file=%s sheets=%d chars=%d",
                     file_path, len(wb.sheetnames), len(result))
        return result
